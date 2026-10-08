#!/usr/bin/env python3
"""Gera peça processual em .docx no timbrado do Lazzari & Ghidorsi Advogados.

Uso:
    python3 scripts/gerar_peca.py peca.md --saida /mnt/user-data/outputs/peca.docx
    python3 scripts/gerar_peca.py peca.md --verificar      # só a revisão

O texto de entrada é um markdown simples, com estas marcações:

    [CENTRO] texto           parágrafo centralizado (endereçamento)
    [AUTOS] texto            linha de autos, sem recuo
    # 1. DOS FATOS           título de seção (negrito)
    > texto                  citação recuada 4 cm, fonte 10, espaçamento simples
                             (linhas ">" seguidas formam um único parágrafo;
                             uma linha ">" vazia separa parágrafos da citação)
    | A | B |                tabela (primeira linha é o cabeçalho; a linha
    |---|---|                separadora é ignorada)
    [ASSINATURA] NOME | OAB/SC 64.362
                             bloco de assinatura; duas linhas [ASSINATURA]
                             seguidas saem lado a lado
    **negrito**  *itálico*   ênfase dentro de qualquer parágrafo

Parágrafos são separados por linha em branco. Quebras simples de linha dentro
de um parágrafo viram espaço.

Toda geração roda antes a revisão automática (terminologia, expressões
proibidas, negritos, remissões a itens, pendências [CONFIRMAR] e espaços de
jurisprudência) e imprime o relatório. A revisão não substitui a leitura.
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
CANDIDATOS_TIMBRADO = [
    os.path.join(RAIZ, "assets", "timbrado-peca.docx"),
    os.path.join(AQUI, "timbrado-peca.docx"),
]

FONTE = "Century Gothic"
CORPO = Pt(12)
CITACAO = Pt(10)
TABELA = Pt(10)
RECUO = Cm(1.27)
RECUO_CITACAO = Cm(4)
ESPACO = Pt(18)          # equivale à linha em branco entre parágrafos

TERMOS_VEDADOS = ["autor", "autora", "autores", "réu", "ré", "réus", "demandante",
                  "demandado", "demandada", "acionante", "suplicante"]
EXPRESSOES_VEDADAS = ["data venia", "in casu", "hodiernamente", "é cediço",
                      "mister se faz", "resta cristalino", "salta aos olhos",
                      "teratológic", "absurdo", "!"]


# --------------------------------------------------------------------------- #
# leitura do texto marcado
# --------------------------------------------------------------------------- #

def blocos(texto):
    """Converte o texto marcado numa lista de blocos (tipo, conteúdo)."""
    linhas = texto.replace("\r\n", "\n").split("\n")
    saida, i = [], 0
    while i < len(linhas):
        ln = linhas[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            saida.append(("titulo", ln[2:].strip()))
            i += 1
        elif ln.startswith("[CENTRO]"):
            saida.append(("centro", ln[len("[CENTRO]"):].strip()))
            i += 1
        elif ln.startswith("[AUTOS]"):
            saida.append(("autos", ln[len("[AUTOS]"):].strip()))
            i += 1
        elif ln.startswith("[ASSINATURA]"):
            grupo = []
            while i < len(linhas) and linhas[i].startswith("[ASSINATURA]"):
                partes = [p.strip() for p in linhas[i][len("[ASSINATURA]"):].split("|")]
                grupo.append(partes)
                i += 1
            saida.append(("assinatura", grupo))
        elif ln.startswith(">"):
            paragrafos, atual = [], []
            while i < len(linhas) and linhas[i].startswith(">"):
                conteudo = linhas[i][1:].strip()
                if conteudo:
                    atual.append(conteudo)
                elif atual:
                    paragrafos.append(" ".join(atual))
                    atual = []
                i += 1
            if atual:
                paragrafos.append(" ".join(atual))
            saida.append(("citacao", paragrafos))
        elif ln.lstrip().startswith("|"):
            linhas_tab = []
            while i < len(linhas) and linhas[i].lstrip().startswith("|"):
                cel = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cel if c):
                    linhas_tab.append(cel)
                i += 1
            saida.append(("tabela", linhas_tab))
        else:
            atual = []
            while (i < len(linhas) and linhas[i].strip()
                   and not re.match(r"(# |\[CENTRO\]|\[AUTOS\]|\[ASSINATURA\]|>|\s*\|)", linhas[i])):
                atual.append(linhas[i].strip())
                i += 1
            saida.append(("paragrafo", " ".join(atual)))
    return saida


TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")


def trechos(texto):
    """Divide o texto em (conteúdo, negrito, itálico)."""
    out = []
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            out.append((parte[2:-2], True, False))
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            out.append((parte[1:-1], False, True))
        else:
            out.append((parte, False, False))
    return out


# --------------------------------------------------------------------------- #
# montagem do documento
# --------------------------------------------------------------------------- #

def fonte_run(run, tamanho, negrito=False, italico=False):
    run.font.name = FONTE
    run.font.size = tamanho
    run.bold = negrito or None
    run.italic = italico or None
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONTE)


def paragrafo(doc_ou_celula, texto, *, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
              recuo_primeira=RECUO, recuo_esq=None, tamanho=CORPO,
              entrelinha=1.5, depois=ESPACO, negrito_total=False):
    p = doc_ou_celula.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = alinhamento
    pf.first_line_indent = recuo_primeira
    if recuo_esq is not None:
        pf.left_indent = recuo_esq
    pf.line_spacing = entrelinha
    pf.space_before = Pt(0)
    pf.space_after = depois
    for conteudo, neg, ita in trechos(texto):
        fonte_run(p.add_run(conteudo), tamanho, negrito_total or neg, ita)
    return p


def sem_bordas(tabela):
    tbl_pr = tabela._element.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "nil")
        bordas.append(b)
    tbl_pr.append(bordas)


def novo_documento():
    for caminho in CANDIDATOS_TIMBRADO:
        if os.path.exists(caminho):
            doc = Document(caminho)
            corpo = doc.element.body
            for el in list(corpo):
                if el.tag != qn("w:sectPr"):
                    corpo.remove(el)
            return doc
    sys.stderr.write("AVISO: timbrado-peca.docx não encontrado; peça gerada sem timbrado.\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.92)
    return doc


FECHO = re.compile(r"^(Nestes termos|Termos em que|[A-ZÀ-Ú][\wÀ-ú ]+/[A-Z]{2}, \d{1,2}º? de )")
LARGURA_UTIL = Cm(15.9)


def nao_quebrar(tabela):
    """Tabela curta não se parte: nenhuma linha divide e cada linha segura a
    seguinte, de modo que o cabeçalho nunca fica sozinho no pé da página."""
    linhas = tabela.rows
    for k, linha in enumerate(linhas):
        tr_pr = linha._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if k == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        if k < len(linhas) - 1:
            for cel in linha.cells:
                for par_ in cel.paragraphs:
                    par_.paragraph_format.keep_with_next = True


def larguras(tabela, ncol):
    pesos = [1.0] if ncol == 1 else [0.7] + [0.3 / (ncol - 1)] * (ncol - 1)
    tabela.autofit = False
    medidas = [int(LARGURA_UTIL * w) for w in pesos]
    grade = tabela._tbl.tblGrid
    for c, col in enumerate(grade.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(int(medidas[c] / 635)))   # EMU -> twips
    for linha in tabela.rows:
        for c, cel in enumerate(linha.cells):
            cel.width = medidas[c]


def montar(lista, destino):
    doc = novo_documento()
    for tipo, conteudo in lista:
        if tipo == "centro":
            paragrafo(doc, conteudo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                      recuo_primeira=Cm(0), depois=Pt(36))
        elif tipo == "autos":
            paragrafo(doc, conteudo, recuo_primeira=Cm(0), depois=Pt(36))
        elif tipo == "titulo":
            paragrafo(doc, conteudo, negrito_total=True)
        elif tipo == "paragrafo":
            p = paragrafo(doc, conteudo)
            if FECHO.match(conteudo):
                # fecho, data e assinatura ficam juntos na mesma página
                p.paragraph_format.keep_with_next = True
                if not conteudo.startswith(("Nestes", "Termos")):
                    p.paragraph_format.space_after = Pt(48)
        elif tipo == "citacao":
            for n, texto in enumerate(conteudo):
                ultimo = n == len(conteudo) - 1
                paragrafo(doc, texto, recuo_primeira=Cm(0), recuo_esq=RECUO_CITACAO,
                          tamanho=CITACAO, entrelinha=1.0,
                          depois=ESPACO if ultimo else Pt(6))
        elif tipo == "tabela":
            if not conteudo:
                continue
            if doc.paragraphs:
                # a frase que apresenta a tabela não fica sozinha no pé da página
                doc.paragraphs[-1].paragraph_format.keep_with_next = True
            ncol = max(len(l) for l in conteudo)
            tab = doc.add_table(rows=0, cols=ncol)
            tab.style = "Table Grid"
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            for n, linha in enumerate(conteudo):
                celulas = tab.add_row().cells
                for c in range(ncol):
                    texto = linha[c] if c < len(linha) else ""
                    cel = celulas[c]
                    cel.paragraphs[0]._element.getparent().remove(cel.paragraphs[0]._element)
                    numerica = c > 0 and re.match(r"^\(?R\$", texto.replace("*", "").strip())
                    paragrafo(cel, texto, recuo_primeira=Cm(0), tamanho=TABELA,
                              entrelinha=1.0, depois=Pt(2),
                              alinhamento=(WD_ALIGN_PARAGRAPH.RIGHT if numerica
                                           else WD_ALIGN_PARAGRAPH.LEFT),
                              negrito_total=(n == 0))
            nao_quebrar(tab)
            larguras(tab, ncol)
            paragrafo(doc, "", depois=Pt(6))
        elif tipo == "assinatura":
            if len(conteudo) == 1:
                for n, linha in enumerate(conteudo[0]):
                    paragrafo(doc, linha, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                              recuo_primeira=Cm(0), entrelinha=1.0, depois=Pt(0),
                              negrito_total=(n == 0))
            else:
                tab = doc.add_table(rows=1, cols=len(conteudo))
                sem_bordas(tab)
                for c, bloco in enumerate(conteudo):
                    cel = tab.rows[0].cells[c]
                    cel.paragraphs[0]._element.getparent().remove(cel.paragraphs[0]._element)
                    for n, linha in enumerate(bloco):
                        paragrafo(cel, linha, alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
                                  recuo_primeira=Cm(0), entrelinha=1.0, depois=Pt(0),
                                  negrito_total=(n == 0))
    doc.core_properties.author = "Lazzari & Ghidorsi Advogados"
    doc.save(destino)


# --------------------------------------------------------------------------- #
# revisão automática
# --------------------------------------------------------------------------- #

def revisar(lista):
    corpo = [c for t, c in lista if t == "paragrafo"]
    titulos = [c for t, c in lista if t == "titulo"]
    texto_corpo = "\n".join(corpo)
    texto_todo = "\n".join(
        c if isinstance(c, str) else " ".join(" ".join(x) if isinstance(x, list) else x for x in c)
        for _, c in lista)
    achados = []

    for termo in TERMOS_VEDADOS:
        n = len(re.findall(rf"(?<![\wÀ-ú]){termo}(?![\wÀ-ú])", texto_corpo, re.I))
        if n:
            achados.append(f"Terminologia: '{termo}' aparece {n}x no corpo (usar o par processual).")
    for expr in EXPRESSOES_VEDADAS:
        if expr == "!":
            n = texto_corpo.count("!")
        else:
            n = len(re.findall(re.escape(expr), texto_corpo, re.I))
        if n:
            achados.append(f"Estilo: expressão vedada '{expr}' aparece {n}x.")
    if "—" in texto_corpo or " – " in texto_corpo:
        achados.append("Estilo: há travessão no corpo; reestruturar a frase.")

    preambulo_idx = next((k for k, (t, _) in enumerate(lista) if t == "paragrafo"), None)
    negritos = sum(len(re.findall(r"\*\*.+?\*\*", c))
                   for k, (t, c) in enumerate(lista) if t == "paragrafo" and k != preambulo_idx)
    if negritos > 5:
        achados.append(f"Ênfase: {negritos} trechos em negrito no corpo (máximo 5).")

    numeros = set()
    for t in titulos:
        m = re.match(r"(\d+(?:\.\d+)*)\.?\s", t)
        if m:
            numeros.add(m.group(1))
    for ref in re.findall(r"\bite(?:m|ns)\s+((?:\d+(?:\.\d+)*)(?:(?:,\s*|\s+e\s+|\s+a\s+)\d+(?:\.\d+)*)*)",
                          texto_todo):
        for numero in re.findall(r"\d+(?:\.\d+)*", ref):
            if numero not in numeros:
                achados.append(f"Remissão: 'item {numero}' não corresponde a nenhum título da peça.")

    confirmar = re.findall(r"\[CONFIRMAR:[^\]]*\]", texto_todo)
    juris = re.findall(r"\[JURISPRUDÊNCIA A INSERIR[^\]]*\]", texto_todo)
    if re.search(r"\d{1,3}(?:\.\d{3})+\.\d{2}\b", texto_todo):
        achados.append("Números: valor com ponto no lugar da vírgula decimal (ex.: 783.490.20).")
    return achados, confirmar, juris


def imprimir_relatorio(achados, confirmar, juris):
    print("REVISÃO AUTOMÁTICA")
    if achados:
        for a in sorted(set(achados)):
            print(f"  [ATENÇÃO] {a}")
    else:
        print("  Nenhum problema de terminologia, estilo, negrito ou remissão detectado.")
    print(f"  Pendências [CONFIRMAR]: {len(confirmar)}")
    for c in confirmar:
        print(f"    - {c}")
    print(f"  Espaços de jurisprudência: {len(juris)}")
    for j in juris:
        print(f"    - {j}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("entrada", help="arquivo .md com a peça marcada")
    ap.add_argument("--saida", help="caminho do .docx a gerar")
    ap.add_argument("--verificar", action="store_true", help="só roda a revisão, sem gerar o .docx")
    args = ap.parse_args()

    with open(args.entrada, encoding="utf-8") as f:
        lista = blocos(f.read())
    imprimir_relatorio(*revisar(lista))
    if args.verificar:
        return
    destino = args.saida or os.path.splitext(args.entrada)[0] + ".docx"
    montar(lista, destino)
    print(f"\nPeça gerada: {destino}")


if __name__ == "__main__":
    main()
